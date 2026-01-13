import io, zipfile, tarfile, traceback
import base64
import fitz  # PyMuPDF
import docx, pptx, openpyxl
from PIL import Image
import pytesseract
from typing import Tuple, Dict, Any, List

def guess_mime(filename: str) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):   return "application/pdf"
    if name.endswith(".docx"):  return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if name.endswith(".pptx"):  return "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    if name.endswith(".xlsx"):  return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if name.endswith(".zip"):   return "application/zip"
    if name.endswith(".tar"):   return "application/x-tar"
    if name.endswith(".tar.gz") or name.endswith(".tgz"): return "application/x-gtar"  # tar+gz
    if name.endswith(".gz"):    return "application/gzip"  # single gzip stream
    if name.endswith(".txt"):   return "text/plain"
    if name.endswith(".png"):   return "image/png"
    if name.endswith(".jpg") or name.endswith(".jpeg"): return "image/jpeg"
    if name.endswith(".gif"):   return "image/gif"
    if name.endswith(".bmp"):   return "image/bmp"
    if name.endswith(".tif") or name.endswith(".tiff"): return "image/tiff"
    return "application/octet-stream"

def extract_any(file_bytes: bytes, filename: str) -> Tuple[str, Dict[str, Any]]:
    """
    Returns a dict in meta with 'text', 'images', 'files' so callers can index and attach images.
    For backward compatibility with your callers, the function still *returns text as first value*,
    and packs images/files inside meta.
    """
    meta: Dict[str, Any] = {"filename": filename}
    mime = guess_mime(filename)
    meta["mime"] = mime

    text_parts: List[str] = []
    images: List[Dict[str, Any]] = []
    files: List[Dict[str, Any]] = []

    try:
        # TEXT
        if mime.startswith("text/"):
            text = file_bytes.decode("utf-8", errors="ignore")
            text_parts.append(text)

        # PDF
        elif mime == "application/pdf":
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            text_parts.append(" ".join(page.get_text() for page in doc))

        # DOCX
        elif mime.endswith("wordprocessingml.document"):
            docx_file = docx.Document(io.BytesIO(file_bytes))
            text_parts.append("\n".join(p.text for p in docx_file.paragraphs))

        # PPTX
        elif mime.endswith("presentationml.presentation"):
            pres = pptx.Presentation(io.BytesIO(file_bytes))
            texts = [shape.text for slide in pres.slides for shape in slide.shapes if hasattr(shape, "text")]
            text_parts.append("\n".join(texts))

        # XLSX
        elif mime.endswith("spreadsheetml.sheet"):
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            rows = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    rows.append(" ".join("" if c is None else str(c) for c in row))
            text_parts.append("\n".join(rows))

        # IMAGES: Keep the raw bytes AND (optionally) OCR text
        elif mime.startswith("image/"):
            # keep original image
            images.append({
                "source": filename,
                "mime": mime,
                "content": file_bytes,   # raw bytes (for vision model)
            })
            # optional OCR text
            try:
                img = Image.open(io.BytesIO(file_bytes))
                ocr_text = pytesseract.image_to_string(img)
                if ocr_text and ocr_text.strip():
                    text_parts.append(ocr_text)
            except Exception as ocr_err:
                # OCR failure shouldn't block the image ingestion
                text_parts.append("")  # keep consistent return

        # ZIP: recurse, collect texts + images + files
        elif mime == "application/zip":
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
                for name in z.namelist():
                    # skip directories
                    if name.endswith("/"):
                        continue
                    with z.open(name) as f:
                        data = f.read()
                        t, child_meta = extract_any(data, name)
                        if t:
                            text_parts.append(f"\n# {name}\n{t}")
                        # propagate child assets
                        for img in child_meta.get("images", []):
                            images.append(img)
                        for fil in child_meta.get("files", []):
                            files.append(fil)

        # TAR / TGZ
        elif mime in ["application/x-tar", "application/x-gtar"]:
            with tarfile.open(fileobj=io.BytesIO(file_bytes), mode="r:*") as tar:
                for m in tar.getmembers():
                    if not m.isfile():
                        continue
                    f = tar.extractfile(m)
                    if f:
                        data = f.read()
                        t, child_meta = extract_any(data, m.name)
                        if t:
                            text_parts.append(f"\n# {m.name}\n{t}")
                        for img in child_meta.get("images", []):
                            images.append(img)
                        for fil in child_meta.get("files", []):
                            files.append(fil)

        # Single GZIP stream (not a tarball): we’ll just surface as a “file” blob
        elif mime == "application/gzip":
            files.append({
                "source": filename,
                "mime": mime,
                "content": file_bytes
            })

        else:
            # Unknown binary -> surface as file, avoid losing data
            files.append({
                "source": filename,
                "mime": mime,
                "content": file_bytes
            })

    except Exception as e:
        print(f"Extraction failed for {filename}: {e}")
        traceback.print_exc()

    # finalize
    meta["text"] = "\n".join(tp for tp in text_parts if isinstance(tp, str))
    meta["images"] = images
    meta["files"] = files

    # For backward-compatibility with your current callers
    return meta["text"], meta
